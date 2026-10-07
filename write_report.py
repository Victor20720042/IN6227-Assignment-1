"""Fill the assignment report following IN6227-Reports-Template.doc."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
FIG = ROOT / "outputs" / "figures"
OUT = ROOT / "IN6227-Assignment-1-Report.docx"

NAVY = RGBColor(0x1F, 0x4E, 0x79)


def set_run_font(run, name="Times New Roman", size=10, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_text(p, text, **kwargs):
    run = p.add_run(text)
    set_run_font(run, **kwargs)
    return run


def set_paragraph(p, *, after=6, before=0, line=1.08, align="justify", first_line=0):
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = line
    p.paragraph_format.first_line_indent = Cm(first_line)
    align_map = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    p.alignment = align_map[align]


def heading(doc, text):
    p = doc.add_paragraph()
    set_paragraph(p, after=3, before=6, line=1.08, align="left")
    add_text(p, text.upper(), size=11, bold=True, color=NAVY)
    return p


def body(doc, text, first_line=0.5):
    p = doc.add_paragraph()
    set_paragraph(p, after=4, before=0, line=1.08, align="justify", first_line=first_line)
    add_text(p, text, size=10)
    return p


def caption(doc, text):
    p = doc.add_paragraph()
    set_paragraph(p, after=8, before=2, line=1.0, align="center")
    add_text(p, text, size=9, italic=True)
    return p


def shade_cell(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "8FAADC")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def write_cell(cell, text, *, bold=False, size=9, align="center", color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    set_paragraph(p, after=0, before=0, line=1.0, align=align)
    add_text(p, text, size=size, bold=bold, color=color)
    set_cell_border(cell)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def build():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.clear()
    set_paragraph(hp, after=0, align="right", line=1.0)
    add_text(hp, "IN6227 Data Mining  •  Assignment 1  •  Variant-1", size=9, italic=True, color=NAVY)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    set_paragraph(fp, after=0, align="center", line=1.0)
    add_text(fp, "Wee Kim Wee School of Communication and Information  •  Page ", size=8, italic=True)
    # PAGE field
    run = fp.add_run()
    set_run_font(run, size=8, italic=True)
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    add_text(fp, " of 2", size=8, italic=True)

    title = doc.add_paragraph()
    set_paragraph(title, after=2, before=0, line=1.15, align="center")
    add_text(
        title,
        "Decision Trees versus Naive Bayes: A Course-Grounded Classification Workflow",
        size=14,
        bold=True,
        color=NAVY,
    )

    author = doc.add_paragraph()
    set_paragraph(author, after=0, align="center", line=1.0)
    add_text(author, "Author Name, Matric Number", size=11, bold=True)

    meta = doc.add_paragraph()
    set_paragraph(meta, after=2, align="center", line=1.0)
    add_text(meta, "IN6227-Assignment-1    Variant-1", size=10)

    git = doc.add_paragraph()
    set_paragraph(git, after=8, align="center", line=1.0)
    add_text(git, "Code repository: ", size=9, italic=True)
    add_text(git, "[insert GitHub URL]", size=9, italic=True, color=NAVY)

    heading(doc, "Introduction")
    body(
        doc,
        "This report runs the supplied binary classification task through the IN6227 pipeline "
        "(Lectures 1–5). After deleting four unlabelled rows, the files contain 31,109 training "
        "and 13,333 test records, 15 attributes, and a yes/no label. The positive class is the "
        "minority (P(yes) ≈ 0.24), so always predicting no already scores 76.2% accuracy. Names "
        "are semantically masked, so decisions rest on the data. We compare a CART tree with "
        "naive Bayes: both are taught in Lectures 4–5, they need different preprocessing, and "
        "they disagree about attribute dependence. The grade target is a justified workflow, "
        "not peak accuracy.",
    )

    heading(doc, "Methods or Procedures")
    body(
        doc,
        "Exploration used Lecture 3 summaries and plots against Lecture 2’s quality issues. "
        "Explicit missingness is 1–7 cells per column; four unlabelled rows were dropped; there "
        "are no exact duplicates. Unknown occupies 5.68% of personal_interest and 5.66% of "
        "species. 264 rows have performance_score = stability_index = 0, all labelled no: the "
        "zeros look coded, but the block is class-pure, so imputing it would erase a pattern. "
        "load_ratio and activity_duration have long right tails. Class balance matches across "
        "the given split. Statistics estimated from data are fitted inside cross-validation "
        "folds. Numeric gaps use the median; nominal gaps become an explicit Missing level. "
        "Levels with fewer than 50 training rows, and unseen test levels, share one reserved "
        "code (Lecture 4’s multi-way split bias). Neither model uses distances, so attributes "
        "are not scaled. Sensitivity runs that kept the same training rows moved cross-validated "
        "F1 by less than the fold standard deviation (0.005 tree, 0.006 naive Bayes); defaults "
        "were kept. No features were constructed. Three numerics remain collinear (r = 0.89, "
        "0.88, 0.83), violating naive Bayes’ independence assumption; they were retained so that "
        "violation could be tested. Imbalance is handled by class_weight = balanced on the tree "
        "(Lecture 4 cost matrix), not by resampling.",
    )
    body(
        doc,
        "Hyperparameters were chosen by stratified five-fold cross-validation on the training "
        "file, maximising positive-class F1; the test file was scored once. The selected tree "
        "uses Gini, max_depth = 10, min_samples_leaf = 100, cost-complexity α = 0 and balanced "
        "weights (depth and leaf size are pre-pruning; α is CART post-pruning). Naive Bayes used "
        "Lecture 5’s first option for continuous attributes, with Lecture 2’s equal-width, "
        "equal-frequency and k-means bins: 10 k-means bins and Laplace α = 0.1. A "
        "Gaussian-plus-categorical variant was indistinguishable from this model and is omitted.",
    )

    heading(doc, "Results")
    body(
        doc,
        "Table 1 places both models beside two trivial baselines. Accuracy barely moves (tree "
        "77.8%, naive Bayes 80.1%, always-no 76.2%). F1, recall and ranking do. The majority "
        "rule detects none of 3,168 positives (F1 = 0; ROC-AUC = 0.50; PR-AUC = 0.24, the base "
        "rate). The tree recovers 83.3% of them (TP 2,638, FN 530, FP 2,431) and naive Bayes "
        "79.6% (TP 2,521, FN 647, FP 2,013). A depth sweep reproduces Lecture 4’s "
        "under/overfitting sketch: training error falls to zero, cross-validated error bottoms "
        "near depth 6 (0.171) and rises to 0.227 unpruned. Depth 10 was selected because the "
        "tuning criterion was F1, not error rate.",
    )

    table = doc.add_table(rows=5, cols=7)
    table.autofit = True
    headers = ["Model", "Accuracy", "Precision", "Recall", "F1", "ROC-AUC", "PR-AUC"]
    rows = [
        ["Always no (baseline)", "0.762", "0.000", "0.000", "0.000", "0.500", "0.238"],
        ["Stratified random", "0.633", "0.225", "0.223", "0.224", "0.492", "0.235"],
        ["Decision tree (CART)", "0.778", "0.520", "0.833", "0.641", "0.876", "0.670"],
        ["Naive Bayes (binned)", "0.801", "0.556", "0.796", "0.655", "0.881", "0.686"],
    ]
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        write_cell(cell, h, bold=True, size=8, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(cell, "1F4E79")
    for i, row in enumerate(rows):
        prevent_row_split(table.rows[i + 1])
        fill = "D6E3F0" if i % 2 == 0 else "FFFFFF"
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            write_cell(cell, val, bold=(j == 0), size=8, align=("left" if j == 0 else "center"))
            shade_cell(cell, fill)
    prevent_row_split(table.rows[0])
    caption(doc, "Table 1. Held-out test metrics (yes = positive). 95% accuracy CIs: tree [0.771, 0.785]; naive Bayes [0.794, 0.807].")

    fig_p = doc.add_paragraph()
    set_paragraph(fig_p, after=0, align="center", line=1.0)
    fig_p.add_run().add_picture(str(FIG / "results_roc_pr.png"), width=Cm(14.2))
    caption(doc, "Figure 1. ROC (left) and precision–recall (right) on the test set. Dashed PR line: base rate 0.24.")

    heading(doc, "Discussion")
    body(
        doc,
        "The models classify; accuracy hides it. Naive Bayes records 2,521 true positives at the "
        "cost of 2,013 false alarms—about 508 extra correct rows over always-no, or +3.8 accuracy "
        "points. Accuracy prices a missed yes and a false alarm equally, which is why a 24% "
        "positive class makes detection look unprofitable (Lecture 4’s 99.9%-accurate classifier "
        "that finds no minority cases). F1, PR-AUC and a cost matrix reverse the reading: at 1:1 "
        "cost the majority rule is already dearer (3,168 vs 2,660); at 5:1 the totals are 15,840 "
        "vs 5,248. The 95% interval for the always-no vs naive-Bayes error gap is [0.028, 0.048]; "
        "for the two learned models it is [0.013, 0.032]. Both assume independent test sets; here "
        "the models share one file, so the test is conservative.",
    )
    body(
        doc,
        "Figure 1’s ROC curves almost coincide. The accuracy gap is mostly a threshold effect: "
        "balanced weights push the tree to recall 0.833 and precision 0.520; at F1-maximising "
        "thresholds the two F1 scores differ by less than 0.003. Dropping two collinear numerics "
        "raised naive Bayes’ cross-validated F1 by 0.005 (inside fold noise) and slightly lowered "
        "test F1: the direction matches the independence assumption, the effect is not shown. "
        "Most signal is nominal (P(yes) across mineral_type spans 0.73). Precision 0.52–0.56 means "
        "a predicted yes is wrong nearly half the time—screens, not verdicts. Two unrelated "
        "inducers stopping at the same ranking quality points to the attributes, not the "
        "algorithm. Deleting Unknown rows raised cross-validated F1 while lowering test F1, "
        "because validation then used an easier population; the untouched test file is the "
        "comparable yardstick.",
    )

    heading(doc, "Conclusion")
    body(
        doc,
        "A lecture-grounded workflow—explore, clean with named missing-value strategies, "
        "preprocess per model family, tune by stratified cross-validation, and evaluate with "
        "more than accuracy—shows that a pruned tree and a discretised naive Bayes each recover "
        "about four fifths of the minority class. Accuracy rises only a few points over always "
        "predicting no, which is a property of the metric under imbalance. Precision is modest "
        "and naive Bayes’ independence assumption is violated; a more elaborate inducer is "
        "unlikely to remove either limit on these attributes.",
    )

    heading(doc, "References")
    refs = [
        "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., Cournapeau, D., Brucher, M., Perrot, M., & Duchesnay, É. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
        "Tan, P.-N., Steinbach, M., Karpatne, A., & Kumar, V. (2019). Introduction to data mining (2nd ed.). Pearson.",
    ]
    for i, ref in enumerate(refs):
        p = doc.add_paragraph()
        indent = Cm(0.75)
        p.paragraph_format.left_indent = indent
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.space_after = Pt(3 if i < len(refs) - 1 else 0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.08
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_text(p, ref, size=9)

    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    build()
